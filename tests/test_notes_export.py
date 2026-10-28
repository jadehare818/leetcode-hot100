import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import app as tracker


class NotesExportTests(unittest.TestCase):
    def setUp(self):
        self.problems = [
            {"id": 49, "title": "字母异位词分组", "category": "哈希", "slug": "group-anagrams"},
            {"id": 1, "title": "两数之和", "category": "哈希", "slug": "two-sum"},
        ]

    def workbook(self, progress, problems=None):
        stream = tracker.build_notes_workbook(
            self.problems if problems is None else problems, progress
        )
        workbook = load_workbook(stream)
        self.addCleanup(workbook.close)
        return workbook.active

    def test_all_problems_in_catalog_order_and_missing_notes(self):
        sheet = self.workbook({"49": {"note": "分类\n使用哈希", "cheatsheet": "```python\nprint('中文')\n```", "status": "solid"}, "999": {"note": "orphan"}})
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.max_column, 6)
        self.assertEqual([sheet.cell(row, 1).value for row in (2, 3)], [49, 1])
        self.assertEqual(sheet["D2"].value, "分类\n使用哈希")
        self.assertEqual(sheet["E2"].value, "```python\nprint('中文')\n```")
        self.assertIsNone(sheet["D3"].value)
        self.assertIsNone(sheet["E3"].value)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:F3")
        self.assertTrue(sheet["D2"].alignment.wrap_text)
        self.assertEqual(sheet["F2"].hyperlink.target, "https://leetcode.cn/problems/group-anagrams/")

    def test_formula_like_notes_remain_text_and_controls_are_removed(self):
        sheet = self.workbook({"49": {"note": '=HYPERLINK("bad")\n中文\x00', "cheatsheet": "#REF!"}})
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["D2"].value, '=HYPERLINK("bad")\n中文')
        self.assertEqual(sheet["E2"].data_type, "s")

    def test_long_notes_are_preserved_in_continuation_columns(self):
        note = "中" * 70000
        cheat = "x" * 40000
        sheet = self.workbook({"49": {"note": note}, "1": {"cheatsheet": cheat}})
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["D2"].value + sheet["G2"].value + sheet["H2"].value, note)
        self.assertEqual(sheet["E3"].value + sheet["I3"].value, cheat)
        self.assertEqual(sheet.auto_filter.ref, "A1:I3")

    def test_empty_progress_and_empty_catalog(self):
        self.assertEqual(self.workbook({}).max_row, 3)
        sheet = self.workbook({}, problems=[])
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.max_column, 6)

    def test_download_uses_local_notes_without_writing_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            catalog = root / "problems.json"
            seed = root / "progress.json"
            local = root / "progress.local.json"
            config = root / "config.local.json"
            catalog.write_text(json.dumps({"categories": [{"name": "哈希", "problems": self.problems}]}), encoding="utf-8")
            seed.write_text('{"49": {"note": "seed"}}', encoding="utf-8")
            local.write_text('{"49": {"note": "local"}}', encoding="utf-8")
            config.write_text('{}', encoding="utf-8")
            before = {path.name: hashlib.sha256(path.read_bytes()).digest() for path in root.iterdir()}
            with patch.multiple(tracker, PROBLEMS_PATH=catalog, PROGRESS_PATH=seed, PROGRESS_LOCAL_PATH=local):
                response = tracker.app.test_client().get("/api/export/notes.xlsx")
                self.assertEqual(response.status_code, 200)
                self.assertIn("attachment;", response.headers["Content-Disposition"])
                self.assertIn(".xlsx", response.headers["Content-Disposition"])
                self.assertEqual(response.headers["Cache-Control"], "no-store")
                workbook = load_workbook(BytesIO(response.data))
                self.assertEqual(workbook.active["D2"].value, "local")
                workbook.close()
                response.close()
            after = {path.name: hashlib.sha256(path.read_bytes()).digest() for path in root.iterdir()}
            self.assertEqual(before, after)

    def test_both_pages_link_to_download(self):
        client = tracker.app.test_client()
        for url in ("/problems", "/cheatsheet"):
            with self.subTest(url=url):
                response = client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertIn(b'href="/api/export/notes.xlsx"', response.data)


if __name__ == "__main__":
    unittest.main()
